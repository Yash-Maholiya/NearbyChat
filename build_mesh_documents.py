from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT = Path('project_documents')
OUT.mkdir(exist_ok=True)
NAVY = '17365D'
BLUE = 'D9EAF7'
PALE = 'F4F8FB'
GRAY = 'D9D9D9'

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), color); tcPr.append(shd)

def border(cell):
    tcPr = cell._tc.get_or_add_tcPr(); b = tcPr.first_child_found_in('w:tcBorders')
    if b is None: b = OxmlElement('w:tcBorders'); tcPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement('w:' + edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),GRAY); b.append(e)

def cell_text(cell, text, bold=False, color=None, size=9.5):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)

def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=OxmlElement('w:'+m); node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa'); mar.append(node)

def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,NAVY); cell_text(c,h,True,'FFFFFF',9.5); set_cell_margins(c)
        if widths: c.width=Inches(widths[i])
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            if ri%2: shade(cells[i], PALE)
            cell_text(cells[i],val,size=9); set_cell_margins(cells[i])
            if widths: cells[i].width=Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

def configure(doc):
    sec=doc.sections[0]; sec.top_margin=Inches(.75); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(10.5); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.08
    for n,size in [('Title',25),('Subtitle',13),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
        s=styles[n]; s.font.name='Aptos Display' if n!='Normal' else 'Aptos'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
        if n.startswith('Heading'): s.font.bold=True; s.paragraph_format.space_before=Pt(14); s.paragraph_format.space_after=Pt(6); s.paragraph_format.keep_with_next=True
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=footer.add_run('Secure Peer to Peer Mesh Communication System'); r.font.size=Pt(8); r.font.color.rgb=RGBColor(100,100,100)

def title_page(doc, title, subtitle):
    p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(120); p.add_run(title)
    p=doc.add_paragraph(style='Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(subtitle)
    doc.add_paragraph('A software only wireless mesh network prototype using Mininet WiFi and Python', style='Subtitle').alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('\nPrepared for academic project submission').alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Prepared by: ____________________\nInstitution: ____________________\nDate: ____________________').alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

def h(doc, text, level=1): doc.add_heading(text, level=level)
def p(doc, text, boldlead=None):
    para=doc.add_paragraph()
    if boldlead:
        r=para.add_run(boldlead); r.bold=True
    para.add_run(text)
    return para
def bullets(doc, items):
    for x in items: doc.add_paragraph(x, style='List Bullet')
def steps(doc, items):
    for x in items: doc.add_paragraph(x, style='List Number')
def codeblock(doc, text):
    para=doc.add_paragraph(); para.paragraph_format.left_indent=Inches(.25); para.paragraph_format.space_after=Pt(8)
    for line in text.splitlines():
        r=para.add_run(line+'\n'); r.font.name='Consolas'; r.font.size=Pt(8.5)

refs=[
['1','Briar Project','How it works','https://briarproject.org/how-it-works/','Accessed September 2026'],
['2','Meshtastic','Off Grid Communication for Everyone','https://meshtastic.org/','Accessed September 2026'],
['3','Open Mesh','B A T M A N advanced source repository','https://github.com/open-mesh-mirror/batman-adv','Accessed September 2026'],
['4','Reticulum Network Stack','What is Reticulum','https://reticulum.network/manual/whatis.html','Accessed September 2026'],
['5','Yggdrasil Network','End to end encrypted IPv6 networking','https://yggdrasil-network.github.io/','Accessed September 2026'],
['6','Mininet WiFi','Emulation Platform for Software Defined Wireless Networks','https://mininet-wifi.github.io/','Accessed September 2026'],
['7','ns 3','Mesh model documentation','https://www.nsnam.org/docs/release/3.40/models/html-new/mesh-design.html','Accessed September 2026'],
['8','CORE','Common Open Research Emulator Documentation','https://coreemu.github.io/core/index.html','Accessed September 2026'],
]

def add_refs(doc):
    h(doc,'References')
    table(doc,['No','Author or Project','Title','URL','Date'],refs,[.35,1.05,2.25,2.3,1.0])

def full_document():
    doc=Document(); configure(doc)
    title_page(doc,'Design and Security Analysis of a Secure Peer to Peer Mesh Communication System','Complete Project Plan and Implementation Guide')
    h(doc,'Executive Summary')
    p(doc,'This project builds and evaluates a small infrastructure free wireless mesh communication system entirely in software. Four or more virtual nodes discover one another, forward messages across multiple hops, and continue delivery when a valid alternate path exists. The prototype uses Mininet WiFi to emulate the wireless network and Python to implement the application and security layers. It does not require a router, Internet connection, radio hardware, or physical devices.')
    p(doc,'The central design principle is that relays help move packets but cannot read or modify user messages. Every node receives a cryptographic identity. A message is encrypted for its intended recipient, signed by its sender, assigned a unique identifier and expiry time, and acknowledged after delivery. A lightweight trust score records observable forwarding reliability without storing message contents.')
    h(doc,'Project Scope')
    table(doc,['Included','Excluded from this prototype'],[
        ['Virtual peer discovery, multi hop forwarding, dynamic link changes, encrypted messaging, signatures, acknowledgements, replay defence, and performance tests.','Real radio propagation, antenna placement, battery testing, licensed spectrum operation, Internet gateways, and tests against systems not owned or authorised by the project team.'],
        ['Four to ten virtual nodes running on one Linux machine or Ubuntu virtual machine.','Claims that a simulator reproduces every physical wireless condition. The report will state this limitation clearly.'],
    ],[3.25,3.25])
    h(doc,'1 Background and Motivation')
    p(doc,'A mesh network is a decentralised network in which devices communicate directly with nearby peers and may relay traffic on behalf of other peers. Unlike a star topology, there is no compulsory central router. This makes mesh communication useful when central infrastructure is unavailable, damaged, congested, or undesirable. Its decentralisation also introduces risks: a relay may be untrusted, topology can change rapidly, and radio traffic can be captured by anyone within range.')
    p(doc,'This work is deliberately a controlled research prototype. Its purpose is to understand discovery, forwarding, route recovery, confidentiality, authenticity, and measurable performance. The system will be evaluated only inside an isolated virtual lab.')
    h(doc,'2 Objectives')
    bullets(doc,[
      'Create an infrastructure free virtual wireless mesh of at least four nodes.',
      'Discover peers and maintain a time limited neighbour table.',
      'Deliver encrypted application messages over one, two, and three relay hops.',
      'Verify sender identity and detect modified, duplicate, expired, and replayed packets.',
      'Recover after a relay or wireless link is removed.',
      'Collect latency, packet delivery ratio, hop count, throughput, route recovery time, and security event data.',
      'Compare the secure implementation with a controlled baseline that has security disabled.'
    ])
    h(doc,'3 Existing Systems')
    table(doc,['System','Approach','Lesson for this project','Reference'],[
      ['Briar','Serverless messaging can use Bluetooth, Wi Fi, the Internet, or removable media; content is end to end encrypted.','Offline communication needs end to end protection, not merely a local radio link.','[1]'],
      ['Meshtastic','Low power LoRa devices form an off grid mesh and use AES 256 channel encryption.','Shows the value of inexpensive decentralised communication but uses a different radio model.','[2]'],
      ['B A T M A N advanced','Linux Layer 2 mesh forwarding implementation.','Shows practical dynamic forwarding beneath applications.','[3]'],
      ['Reticulum','Cryptography based stack with encrypted multi hop transport, identities, and delivery proofs.','Inspires identity based addressing and authenticated acknowledgement design.','[4]'],
      ['Yggdrasil','Self healing peer to peer encrypted IPv6 overlay.','Illustrates that intermediate peers should not decrypt end user traffic.','[5]'],
    ],[1.0,1.8,2.65,1.05])
    p(doc,'Gap addressed. Existing systems are mature and have broader goals. This academic prototype exposes its security decisions and results in a readable dashboard. Its novelty is the combination of QR style trust enrolment, end to end message protection, replay aware forwarding, and a privacy preserving trust score for relays.')
    h(doc,'4 Problem Statement')
    p(doc,'Design, implement, and evaluate a secure peer to peer wireless mesh communication system in which virtual nodes communicate without a central router or Internet connection. The system shall discover peers, relay messages over multiple hops, protect message confidentiality and integrity from intermediate relays, authenticate enrolled devices, resist simple replay and modification attacks, and report measurable performance and security outcomes.')
    h(doc,'5 Requirements')
    table(doc,['Category','Requirement','Acceptance evidence'],[
      ['Functional','A node discovers reachable neighbours automatically.','Neighbour table contains signed hello messages from adjacent nodes.'],
      ['Functional','Node A sends to Node D through intermediate nodes.','Acknowledged delivery with hop count greater than one.'],
      ['Security','Only the intended destination can decrypt payload text.','Relay capture displays encrypted ciphertext, not plaintext.'],
      ['Security','Forged or altered packets are rejected.','Security log records signature or authentication failure.'],
      ['Security','A previously received packet is not delivered twice.','Replay cache rejects duplicate message ID and sequence number.'],
      ['Reliability','A failed path is avoided when a valid alternative exists.','Measured route recovery experiment.'],
      ['Usability','User can send a message and see delivery status locally.','Local web dashboard screenshot and test log.'],
    ],[.8,3.2,2.5])
    h(doc,'6 Proposed Architecture and Novelty')
    p(doc,'The architecture separates radio emulation, forwarding, secure messaging, and visualisation so that each component can be assessed independently. Mininet WiFi provides isolated virtual stations and controlled range, loss, and mobility. The application runs in each station namespace and uses UDP for control packets and application messages.')
    codeblock(doc,'User dashboard\n    |\nSecure messaging layer: recipient encryption | sender signature | replay cache | acknowledgements\n    |\nTrust aware forwarding layer: discovery | neighbour table | duplicate suppression | TTL | relay score\n    |\nVirtual Wi Fi mesh: Mininet WiFi | mesh or ad hoc links | controlled loss and mobility\n    |\nVirtual nodes: nodeA | nodeB | nodeC | nodeD')
    h(doc,'6 1 Novel Contributions',2)
    table(doc,['Contribution','Design decision','Value'],[
      ['Trusted enrolment','The user verifies a new node public key fingerprint through a QR code or short authentication string before trust is granted.','Prevents simple device impersonation in the project network.'],
      ['End to end sealed payload','Payload is encrypted for the destination public key and authenticated; the sender signs the outer packet.','Relays can forward but cannot read or silently change the message.'],
      ['Replay aware forwarding','Each packet includes sender ID, message ID, sequence number, timestamp, expiry, and a bounded replay cache.','Stops duplicate delivery and limits routing loops.'],
      ['Trust aware route preference','A relay score is updated from acknowledgement success, forwarding delay, drops, and invalid events.','Allows comparison between shortest hop routing and reliability aware routing.'],
      ['Privacy preserving audit trail','Logs retain packet hash and event type, never plaintext message content or private keys.','Supports security analysis without creating a content surveillance log.'],
    ],[1.25,3.1,2.15])
    h(doc,'6 2 Packet Format',2)
    table(doc,['Field','Purpose','Security role'],[
      ['Version and type','Identifies hello, data, acknowledgement, or alert packet.','Safe protocol parsing.'],
      ['Sender and destination IDs','Public key fingerprints or short identifiers.','Binds packet to known identities.'],
      ['Message ID and sequence','Unique sender generated values.','Duplicate and replay detection.'],
      ['Timestamp and expiry','Defines validity window.','Rejects stale packets.'],
      ['TTL and hop count','Controls forwarding and records path length.','Limits loops and flooding impact.'],
      ['Encrypted payload','Recipient only plaintext message.','Confidentiality and tamper detection.'],
      ['Sender signature','Signature over canonical packet header and ciphertext.','Authenticity and integrity.'],
    ],[1.35,2.5,2.65])
    h(doc,'7 Threat Model and Controls')
    p(doc,'The prototype assumes an attacker may observe packets, replay previously captured packets, alter packets, introduce an untrusted virtual node, or operate a relay that drops traffic. It does not claim protection from physical device compromise, radio jamming, or global traffic analysis. All security tests must be run only in the project simulator.')
    table(doc,['Threat','Expected effect without defence','Control and expected outcome'],[
      ['Eavesdropping','Relay or observer reads user text.','Authenticated encryption leaves observers with ciphertext only.'],
      ['Modification','Changed packet reaches recipient.','Signature and authenticated decryption reject it.'],
      ['Impersonation','Fake node joins as a valid device.','Trusted public key registry rejects unknown identity.'],
      ['Replay','Old valid message is delivered again.','Message ID, sequence, timestamp, expiry, and cache reject it.'],
      ['Dropper relay','Messages disappear without explanation.','Acknowledgements expose failure; score decreases; alternative route may be selected.'],
      ['Flooding','Packets consume CPU and link capacity.','Rate limit, TTL, payload-size cap, and duplicate suppression.'],
    ],[1.1,2.3,3.1])
    h(doc,'8 Software and Environment')
    table(doc,['Tool','Use in project','Reason'],[
      ['Ubuntu 22 04 or later in VirtualBox or WSL2','Host operating environment.','Linux networking namespaces are required by Mininet WiFi.'],
      ['Mininet WiFi','Virtual stations, Wi Fi modes, range, mobility, and network isolation.','Runs real Linux applications in emulated wireless topology. [6]'],
      ['Python 3.10 or later','Messaging protocol, dashboard, tests, logs, and analysis.','Rapid implementation and good cryptography support.'],
      ['PyNaCl libsodium','Ed25519 signing and authenticated encryption primitives.','Avoids implementing cryptography manually.'],
      ['Flask or FastAPI','Local dashboard API and interface.','Simple node status and test-result display.'],
      ['Wireshark and tcpdump','Authorised packet inspection.','Confirms ciphertext and evaluates protocol traffic.'],
      ['iperf3 and Matplotlib','Baseline transfer test and result graphs.','Repeatable performance measurement.'],
      ['Git','Version control.','Tracks changes, experiments, and reproducibility.'],
    ],[1.4,2.7,2.4])
    h(doc,'9 Step by Step Implementation Plan')
    h(doc,'9 1 Phase 0 Prepare the Host',2)
    steps(doc,[
      'Install Ubuntu in a virtual machine or use a native Linux computer. Allocate at least 4 CPU cores, 8 GB RAM, and 25 GB storage where available.',
      'Create a project folder with directories named topology, node_app, tests, logs, results, docs, and scripts.',
      'Install Mininet WiFi using its official instructions. Verify that the sample Wi Fi topology starts and stops cleanly.',
      'Create a Python virtual environment and install pynacl, flask or fastapi, pytest, pandas, matplotlib, and scapy only if packet construction or inspection is required.',
      'Initialise a Git repository and create a README containing host version, package versions, and execution commands.'
    ])
    h(doc,'9 2 Phase 1 Build the Baseline Topology',2)
    steps(doc,[
      'Write topology/line_topology.py to create nodeA, nodeB, nodeC, and nodeD.',
      'Position nodes so only adjacent stations are within simulated radio range. This forces A to D traffic to use relays.',
      'Enable mesh or ad hoc mode. Assign static addresses only for the virtual experiment; no NAT, router, or Internet gateway is created.',
      'Start a plain UDP echo program on each node and verify one hop, two hop, and three hop reachability.',
      'Capture baseline latency, loss, and throughput. Save raw CSV files in results/baseline.'
    ])
    h(doc,'9 3 Phase 2 Implement Identity and Enrolment',2)
    steps(doc,[
      'On first run, generate an Ed25519 signing key pair and a separate encryption key pair for every node.',
      'Store private keys in a local protected directory. Never write keys to logs or source control.',
      'Derive a short node ID from a hash of the public signing key.',
      'Export public keys and short authentication strings. Simulate QR verification by placing verified public keys in config/trusted_peers.json.',
      'Reject hello messages from identities not in the trusted registry, except during an explicitly enabled enrolment experiment.'
    ])
    h(doc,'9 4 Phase 3 Implement Discovery and Forwarding',2)
    steps(doc,[
      'Broadcast a signed hello packet every few seconds containing node ID, timestamp, supported version, and current status.',
      'Verify incoming hello signatures and expire a neighbour if no valid hello is received before the timeout.',
      'Maintain neighbour_table.json in memory and periodically export safe diagnostic values.',
      'Implement packet duplicate suppression using a bounded cache keyed by sender ID and message ID.',
      'For valid data packets, decrement TTL, increment hop count, and forward only if TTL remains positive and the destination is not local.',
      'Begin with flood forwarding for simplicity, then compare it with a route table or trust-aware preferred-next-hop strategy.'
    ])
    h(doc,'9 5 Phase 4 Implement Secure Messaging',2)
    steps(doc,[
      'Create a canonical packet header so each node signs exactly the same byte sequence.',
      'Encrypt plaintext to the recipient public encryption key using an authenticated encryption primitive. Keep the encrypted payload separate from the signed header.',
      'Sign the header and ciphertext with the sender Ed25519 key.',
      'At the destination, verify the signature before decryption; then validate timestamp, expiry, sender, and replay state before showing the message.',
      'Create a signed acknowledgement that references the delivered message ID. Route the acknowledgement back to the origin.',
      'Mark a message delivered only after a valid acknowledgement is received; otherwise retry a limited number of times.'
    ])
    h(doc,'9 6 Phase 5 Add Trust Score and Dashboard',2)
    steps(doc,[
      'Record only safe metadata: event type, packet hash prefix, time, hop count, acknowledgement result, and peer ID.',
      'Calculate relay score using successful acknowledgement ratio, mean delay, observed drops, invalid forwards, and recent activity. State clearly that this is an experiment heuristic, not proof of malicious intent.',
      'Create a local dashboard displaying topology, peer status, outstanding messages, delivery rate, latency, and security event counts.',
      'Use score only as a route preference after a baseline shortest-path experiment has been completed.'
    ])
    h(doc,'9 7 Phase 6 Test and Evaluate',2)
    steps(doc,[
      'Run each scenario at least three times with the same random seed where applicable.',
      'Test fixed line topology, alternate-path topology, mobility or range change, and relay removal.',
      'Perform controlled authorised security cases: observe ciphertext, submit a modified test packet, replay a previously accepted packet, attempt an unknown identity hello, and configure a relay to drop packets.',
      'Export CSV results and generate charts for latency, delivery ratio, hop count, recovery time, and security overhead.',
      'Write conclusions that distinguish measured facts from limitations and future work.'
    ])
    h(doc,'10 Suggested Directory Structure')
    codeblock(doc,'mesh-project/\n  topology/line_topology.py\n  node_app/main.py\n  node_app/crypto.py\n  node_app/protocol.py\n  node_app/routing.py\n  node_app/dashboard.py\n  config/trusted_peers.json\n  tests/test_crypto.py\n  tests/test_replay.py\n  tests/test_delivery.py\n  results/\n  docs/')
    h(doc,'11 Test Plan')
    table(doc,['Test ID','Scenario','Expected result','Primary metric'],[
      ['T1','A to B direct message','Acknowledged and decrypted by B.','Latency and delivery ratio'],
      ['T2','A to D over B and C','Acknowledged delivery with hop count 3.','Hop count and latency'],
      ['T3','Break B C link; enable alternate route','Path changes or failure is reported honestly.','Route recovery time'],
      ['T4','Observe relay traffic','Payload unreadable to relay.','Packet capture evidence'],
      ['T5','Alter ciphertext or signature','Recipient rejects packet.','Integrity rejection count'],
      ['T6','Replay saved packet','No second user-visible delivery.','Replay rejection count'],
      ['T7','Unknown identity sends hello','Peer does not become trusted.','Unauthorised peer count'],
      ['T8','Relay drops packets','Delivery ratio falls; relay score reflects observable failure.','PDR and score'],
    ],[.55,2.25,2.7,1.4])
    h(doc,'12 Evaluation Method')
    p(doc,'Packet Delivery Ratio equals delivered messages divided by sent messages multiplied by 100. End to end latency is the time between local send and verified acknowledgement. Security overhead is the difference in packet size and measured latency between the secure and baseline modes. Route recovery time begins when a selected link or relay is disabled and ends when a new valid delivery acknowledgement is observed.')
    table(doc,['Experiment','Independent variable','Controlled values','Output'],[
      ['Hop scaling','1, 2, and 3 relay hops','Payload size, trial count, topology seed','Latency and delivery ratio'],
      ['Loss tolerance','0, 2, 5, and 10 percent simulated loss','Nodes, TTL, payload size','Delivery ratio and retries'],
      ['Security cost','Security mode on or off','Topology and message workload','Latency, packet size, CPU estimate'],
      ['Failure recovery','Remove relay or reduce its range','Message interval and retry count','Recovery time'],
      ['Trust preference','Normal relay versus controlled dropper','Alternative route available','Score trend and selected path'],
    ],[1.35,1.55,2.2,1.8])
    h(doc,'13 Timeline')
    table(doc,['Week','Deliverable'],[
      ['1','Literature review, final requirements, host setup, and repository creation.'],
      ['2','Four node baseline topology and direct plus multi hop connectivity.'],
      ['3','Peer discovery, neighbour tables, TTL, and duplicate suppression.'],
      ['4','Keys, trusted registry, signed packets, and encrypted delivery.'],
      ['5','Acknowledgements, retransmission, and route failure handling.'],
      ['6','Trust score and local dashboard.'],
      ['7','Performance tests and controlled security tests.'],
      ['8','Result analysis, report, slides, and demonstration rehearsal.'],
    ],[.7,6.1])
    h(doc,'14 Risks and Mitigation')
    table(doc,['Risk','Mitigation'],[
      ['Mininet WiFi setup issues','Start from official examples; use an Ubuntu VM snapshot before changes; retain a simpler wired namespace fallback for application-layer tests.'],
      ['Overly complex routing','Implement a correct flood plus duplicate-cache baseline before adding route preference.'],
      ['Cryptography misuse','Use PyNaCl high-level APIs; do not write custom ciphers or key-exchange code.'],
      ['Unreproducible tests','Use fixed topology files, saved random seeds, documented commands, and CSV output.'],
      ['Simulation limitations','State them explicitly and avoid claims about real radio range, jamming resistance, or battery life.'],
    ],[1.6,5.2])
    h(doc,'15 Final Deliverables')
    bullets(doc,['Source code and README with reproducible setup instructions.','Mininet WiFi topology scripts.','Secure message protocol specification and trusted public key configuration.','Local dashboard and safe diagnostic logs.','CSV results, graphs, and test evidence.','Final written report, presentation slides, and recorded or live demonstration.'])
    h(doc,'16 Conclusion')
    p(doc,'The software-only approach is sufficient to demonstrate core mesh behaviour and security controls in a repeatable laboratory setting. It removes hardware cost while allowing the project to test discovery, relaying, path disruption, encrypted delivery, authenticity, replay defence, and measured network performance. A future deployment on real devices can reuse the application and security layers after the routing and radio interface are adapted.')
    add_refs(doc)
    doc.save(OUT/'Complete_Project_Plan_Secure_Mesh_System.docx')

def synopsis_document():
    doc=Document(); configure(doc)
    title_page(doc,'Design and Security Analysis of a Secure Peer to Peer Mesh Communication System','Project Synopsis')
    h(doc,'Title')
    p(doc,'Design and Security Analysis of a Secure Peer to Peer Mesh Communication System')
    h(doc,'Introduction')
    p(doc,'This project proposes a software-only prototype of a wireless peer to peer mesh communication system. It will create virtual nodes that communicate directly without a router, Internet connection, or central server. Each node can discover nearby peers, relay a message to the next node, and deliver a message across several hops. For example, if node A cannot reach node D directly, it can send through node B and node C.')
    p(doc,'The project also focuses on security. In a decentralised mesh, the relay that carries a packet may not be trusted. The prototype therefore encrypts user messages for the final recipient, signs messages to prove sender identity, rejects replayed or altered packets, and records safe security events without storing message content. The work will be implemented in a controlled virtual laboratory, making it low-cost and repeatable.')
    h(doc,'Existing Systems')
    table(doc,['System','Summary','Reference'],[
      ['Briar','Secure peer to peer messaging that can use Bluetooth and Wi Fi when the Internet is unavailable. It protects content with end to end encryption.','[1]'],
      ['Meshtastic','Open-source off grid mesh communication over LoRa devices with AES 256 protection.','[2]'],
      ['B A T M A N advanced','A Linux Layer 2 mesh forwarding implementation used to create dynamic mesh links.','[3]'],
      ['Reticulum','A cryptography based network stack providing identities, encrypted multi hop transport, and delivery proofs.','[4]'],
      ['Yggdrasil','A decentralised self healing encrypted peer to peer IPv6 overlay network.','[5]'],
    ],[1.4,4.2,1.15])
    h(doc,'Problem Statement')
    p(doc,'Conventional communication systems depend on routers, cellular networks, or Internet services. When this infrastructure is unavailable, users may be unable to exchange data. A peer to peer mesh network can provide local multi-hop communication, but it raises security concerns because packets are visible to intermediate relay devices. The problem is to build and evaluate a secure infrastructure-free mesh prototype that supports discovery, multi-hop delivery, device authenticity, message integrity, and confidentiality without requiring physical hardware.')
    h(doc,'Proposed System Architecture')
    p(doc,'The proposed system uses Mininet WiFi to emulate four to ten virtual wireless nodes. A Python application runs in every virtual node. The network layer handles peer discovery, TTL controlled forwarding, neighbour status, duplicate suppression, and route selection. The security layer adds public-key identities, destination encryption, sender signatures, timestamps, expiry, and acknowledgement messages.')
    codeblock(doc,'Application dashboard\n  -> secure messaging: encrypt | sign | verify | replay cache | acknowledgement\n  -> forwarding: discovery | neighbours | TTL | duplicate suppression | relay score\n  -> Mininet WiFi: virtual stations | mesh links | loss | mobility\n  -> nodeA <-> nodeB <-> nodeC <-> nodeD')
    h(doc,'Novelty of the Proposed System',2)
    table(doc,['Feature','Novel element'],[
      ['Trusted enrolment','A device is accepted only after its public key fingerprint has been verified through a simulated QR code or short authentication string.'],
      ['End to end secure relaying','Relay nodes forward ciphertext but cannot read the plaintext message. The sender signs the packet and the destination verifies it.'],
      ['Replay aware protocol','Message ID, sequence number, timestamp, expiry time, TTL, and a replay cache prevent duplicate delivery and looped traffic.'],
      ['Trust aware forwarding','Each relay receives an experimental reliability score based on successful delivery acknowledgements and observed packet drops.'],
      ['Privacy preserving audit log','The system stores event type and packet hash prefix rather than message content.'],
    ],[1.6,5.15])
    h(doc,'Tools Used')
    table(doc,['Tool','Purpose'],[
      ['Ubuntu Linux or Ubuntu virtual machine','Host platform for network namespaces and emulation.'],
      ['Mininet WiFi','Creates virtual Wi Fi nodes, mesh links, range, loss, and mobility. [6]'],
      ['Python','Implements peer discovery, forwarding, encryption workflow, tests, and dashboard.'],
      ['PyNaCl libsodium','Provides Ed25519 signatures and authenticated encryption.'],
      ['Flask or FastAPI','Creates a local dashboard for sending messages and viewing status.'],
      ['Wireshark and tcpdump','Captures authorised test traffic to confirm encryption and analyse protocol behaviour.'],
      ['iperf3 Pandas Matplotlib','Measures baseline traffic and prepares result graphs.'],
      ['Git','Tracks source code and experiment changes.'],
    ],[1.7,5.05])
    h(doc,'How the System Will Be Implemented')
    steps(doc,[
      'Install Ubuntu and Mininet WiFi, create a Python environment, and verify a basic virtual wireless topology.',
      'Create four virtual nodes in a line. Configure the simulation so A reaches D only through B and C.',
      'Implement basic peer discovery using signed hello messages and a neighbour table with timeouts.',
      'Implement forwarding with TTL, hop count, duplicate suppression, and bounded retransmission.',
      'Generate public and private keys for each node. Add verified public keys to a trusted registry.',
      'Encrypt each message for its destination and sign the packet with the sender key.',
      'Verify sender signature, expiry, and replay state at the destination before displaying a message.',
      'Return signed acknowledgements and record delivery delay, hops, and success rate.',
      'Add relay reliability score and a local dashboard for topology, messages, events, and test metrics.',
      'Run multi-hop, failure recovery, packet modification, replay, unauthorised-peer, and controlled-dropper experiments.'],)
    h(doc,'Expected Outcome')
    p(doc,'The outcome will be a reproducible virtual mesh-network prototype that proves encrypted multi-hop delivery without a physical router or Internet service. The final evaluation will report latency, packet delivery ratio, hop count, recovery time, security overhead, and the outcome of controlled security tests. The project will also document that simulation does not fully model radio interference, battery life, or physical wireless range.')
    add_refs(doc)
    doc.save(OUT/'Project_Synopsis_Secure_Mesh_System.docx')

if __name__ == '__main__':
    full_document(); synopsis_document()
    print('Created documents in', OUT.resolve())
